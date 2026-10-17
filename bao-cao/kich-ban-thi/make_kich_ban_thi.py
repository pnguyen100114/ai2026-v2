"""Xuất kịch bản thi (Phần 1: Thuyết trình & Bảo vệ) ra file Word để in cầm tay.

Nguồn duy nhất là KICH_BAN_THI.md cùng thư mục. Sửa chữ thì sửa ở đó rồi chạy lại:
    python bao-cao/kich-ban-thi/make_kich_ban_thi.py
Đừng sửa thẳng trong Word, lần chạy sau sẽ ghi đè.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / 'KICH_BAN_THI.md'
DOCX = HERE / 'Kich_ban_thi.docx'

CORAL, NAVY, GREY = 'F2704F', '1B2A4A', '6B7280'
ITEM = re.compile(r'^( *)(- \[[ xX]\] |- |\d+\. )')


# ---------------------------------------------------------------- tách markdown thành khối
def parse(md):
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip() or ln.strip() == '---':
            i += 1
        elif ln.startswith('#'):
            level = len(ln) - len(ln.lstrip('#'))
            blocks.append((('title', 'h1', 'h2')[level - 1], ln[level:].strip())); i += 1
        elif ln.startswith('!['):
            alt, path = re.match(r'!\[(.*)\]\((.*)\)', ln).groups()
            blocks.append(('image', (alt, path))); i += 1
        elif ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            blocks.append(('table', [r for r in rows if not all(set(c) <= set('-: ') for c in r)]))
        elif ln.startswith('>'):
            group = []
            while i < len(lines) and lines[i].startswith('>'):
                group.append(lines[i].lstrip('>').strip()); i += 1
            paras, cur = [], []
            for g in group:                       # dòng `>` trống là chỗ ngắt đoạn trong khối trích
                if g:
                    cur.append(g)
                elif cur:
                    paras.append(cur); cur = []
            if cur:
                paras.append(cur)
            blocks.append(('quote', paras))
        elif ITEM.match(ln):
            items = []                            # (mức lồng, loại, chữ)
            while i < len(lines) and ITEM.match(lines[i]):
                indent, marker = ITEM.match(lines[i]).groups()
                kind = 'check' if marker.startswith('- [') else 'number' if marker[0].isdigit() else 'bullet'
                items.append((len(indent) // 4, kind, marker, lines[i][len(indent) + len(marker):].strip()))
                i += 1
            blocks.append(('list', items))
        else:
            text = []
            while i < len(lines) and lines[i].strip() and not re.match(r'[#>|!]', lines[i]) \
                    and not ITEM.match(lines[i]):
                text.append(lines[i].strip()); i += 1
            blocks.append(('para', ' '.join(text)))
    return blocks


# ---------------------------------------------------------------- dựng file Word
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8), Cm(1.8), Cm(1.8)
WIDTH = Cm(17.2)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'; normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15
for name, size in (('Heading 1', 15), ('Heading 2', 13)):
    s = doc.styles[name]
    s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(NAVY)
    rf = s.element.rPr.rFonts
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        rf.attrib.pop(qn(attr), None)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(attr), 'Times New Roman')
    s.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 10)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True


def runs(p, text, italic=False, size=None, color=None, bold_all=False):
    """Hiểu **đậm** và `mã` trong markdown."""
    for chunk in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not chunk:
            continue
        code = chunk.startswith('`')
        r = p.add_run(chunk[1:-1] if code else chunk.strip('*'))
        r.bold = bold_all or chunk.startswith('**')
        r.italic = italic
        if code:
            r.font.name = 'Consolas'
            r.font.size = Pt((size or 12) - 1.5)
            r.font.color.rgb = RGBColor.from_string(CORAL)
        else:
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def bar(p, color):
    """Vạch màu dọc bên trái đoạn, cho dễ phân biệt lời nói với ghi chú."""
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '8'); left.set(qn('w:color'), color)
    pbdr.append(left)
    p._p.get_or_add_pPr().append(pbdr)


def para(text, **kw):
    align = kw.pop('align', None)
    after = kw.pop('after', None)
    indent = kw.pop('indent', None)
    spacing = kw.pop('spacing', None)
    p = doc.add_paragraph()
    runs(p, text, **kw)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if spacing is not None:
        p.paragraph_format.line_spacing = spacing
    return p


def table(rows):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    size = 10 if len(rows[0]) >= 5 else 11      # bảng nhiều cột thì chữ nhỏ lại cho vừa khổ dọc
    # Cột rộng theo lượng chữ (căn bậc hai để cột dài không nuốt hết), tối thiểu đủ cho tiêu đề cột.
    need = [max(len(r[i]) for r in rows) ** 0.5 for i in range(len(rows[0]))]
    need = [max(w, len(rows[0][i]) ** 0.5 * 1.3, 2.2) for i, w in enumerate(need)]
    widths = [WIDTH * w / sum(need) for w in need]
    for n, row in enumerate(rows):
        tr = t.add_row()
        cant = OxmlElement('w:cantSplit')       # một hàng không bị cắt đôi sang trang sau
        tr._tr.get_or_add_trPr().append(cant)
        cells = tr.cells
        for i, v in enumerate(row):
            cells[i].width = int(widths[i])
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.05
            runs(cells[i].paragraphs[0], v, size=size, bold_all=n == 0)
            if n == 0:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'FDE7E0')
                cells[i]._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


for kind, payload in parse(SRC.read_text(encoding='utf-8')):
    if kind == 'title':
        para(payload, align=WD_ALIGN_PARAGRAPH.CENTER, size=17, color=NAVY, after=4, bold_all=True)
    elif kind == 'h1':
        doc.add_heading(payload, level=1)
    elif kind == 'h2':
        doc.add_heading(payload, level=2)
    elif kind == 'table':
        table(payload)
    elif kind == 'image':
        alt, path = payload
        doc.add_picture(str(HERE / path), width=WIDTH)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    elif kind == 'list':
        for depth, item_kind, marker, text in payload:
            if item_kind == 'check':
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                runs(p, ('☑ ' if 'x' in marker.lower() else '☐ ') + text)
            elif item_kind == 'number':
                # Đánh số tay để giữ đúng số câu (7, 12…) khi danh sách bị ngắt bởi tiêu đề.
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.keep_with_next = True
                runs(p, marker + text)
            else:
                p = doc.add_paragraph(style='List Bullet 2' if depth else 'List Bullet')
            if item_kind == 'bullet':
                runs(p, text)
            p.paragraph_format.space_after = Pt(3)
    elif kind == 'quote':
        # Khối trích là câu phải nói: chữ to, giãn dòng, vạch cam bên trái.
        for lines in payload:
            p = doc.add_paragraph()
            for n, line in enumerate(lines):
                if n:
                    p.add_run().add_break()
                runs(p, line, size=13, color=NAVY)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.3
            bar(p, CORAL)
    else:
        first = payload.startswith('Phần 1:')
        para(payload, after=10 if first else 6, italic=first, color=GREY if first else None,
             align=WD_ALIGN_PARAGRAPH.CENTER if first else None)

doc.save(DOCX)
print('ok', DOCX.relative_to(HERE.parent.parent))
